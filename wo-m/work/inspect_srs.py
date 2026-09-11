from docx import Document
from pathlib import Path

src = Path(r"C:\Users\d1828\Documents\xwechat_files\wxid_i0r32xl0pplw22_3589\msg\file\2026-09\01_智能互动教学平台_Software Requirement Specification_V1.0.docx")
doc = Document(src)
with open(r"C:\Users\d1828\Documents\Codex\2026-09-04\wo-m\work\srs_structure.txt", "w", encoding="utf-8") as f:
    f.write(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}\n")
    for i,p in enumerate(doc.paragraphs):
        if p.text.strip(): f.write(f"P{i}: {p.text}\n")
    for ti,t in enumerate(doc.tables):
        f.write(f"TABLE {ti} rows={len(t.rows)} cols={len(t.columns)}\n")
        for ri,row in enumerate(t.rows):
            f.write(f"R{ri}: " + " | ".join(c.text.replace("\n", " / ") for c in row.cells) + "\n")
print('done')
