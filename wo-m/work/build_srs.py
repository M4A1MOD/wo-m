from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]; DOCS=ROOT/'docs'; OUT=ROOT/'outputs'/'智能互动教学平台_SRS_V1.1_需求评审版.docx'

def font(r,size=10.5,bold=False,color='000000'):
    r.font.name='Microsoft YaHei'; r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); r._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),'Microsoft YaHei'); r.font.size=Pt(size); r.bold=bold; r.font.color.rgb=RGBColor.from_string(color)
def para(d,text='',size=10.5,bold=False,align=None,after=5,style=None):
    p=d.add_paragraph(style=style) if style else d.add_paragraph(); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.18
    if align is not None:p.alignment=align
    font(p.add_run(text),size,bold); return p
def head(d,text,level=1):
    p=d.add_paragraph(style=f'Heading {level}'); p.paragraph_format.space_before=Pt(12 if level==1 else 8); p.paragraph_format.space_after=Pt(5); font(p.add_run(text),15 if level==1 else 12,True); return p
def shade(c,fill):
    shd=c._tc.get_or_add_tcPr().find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); c._tc.get_or_add_tcPr().append(shd)
    shd.set(qn('w:fill'),fill)
def borders(t):
    b=t._tbl.tblPr.first_child_found_in('w:tblBorders')
    if b is None:b=OxmlElement('w:tblBorders'); t._tbl.tblPr.append(b)
    for e in ('top','left','bottom','right','insideH','insideV'):
        x=b.find(qn('w:'+e))
        if x is None:x=OxmlElement('w:'+e); b.append(x)
        x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'6'); x.set(qn('w:color'),'D9D9D9')
def table(d,headers,rows):
    t=d.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; borders(t)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c,'405C9E'); c.text=''; font(c.paragraphs[0].add_run(h),8.8,True,'FFFFFF'); c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
    for ri,row in enumerate(rows):
        cs=t.add_row().cells
        for i,v in enumerate(row):
            cs[i].text='';
            if ri%2:shade(cs[i],'F4F7FC')
            font(cs[i].paragraphs[0].add_run(str(v)),8.5)
    d.add_paragraph().paragraph_format.space_after=Pt(1); return t
def image(d,path,w):
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(path.with_suffix('.png')),width=Inches(w)); return p

d=Document(); s=d.sections[0]; s.top_margin=Inches(.78); s.bottom_margin=Inches(.78); s.left_margin=Inches(.9); s.right_margin=Inches(.9)
d.styles['Normal'].font.name='Microsoft YaHei'; d.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); d.styles['Normal'].font.size=Pt(10.5)
for st in ('Heading 1','Heading 2','Heading 3'): d.styles[st].font.name='Microsoft YaHei'; d.styles[st]._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); d.styles[st].font.color.rgb=RGBColor(0,0,0)

para(d,'智能互动教学平台',27,True,WD_ALIGN_PARAGRAPH.CENTER,20); para(d,'软件需求规格说明书',20,True,WD_ALIGN_PARAGRAPH.CENTER,30); para(d,'需求评审版  V1.1',12,True,WD_ALIGN_PARAGRAPH.CENTER,8); para(d,'项目编号：PRJ-2026-001    文档编号：SRS-2026-001',10.5,False,WD_ALIGN_PARAGRAPH.CENTER,8); para(d,'拟制：项目组    日期：2026-09-07',10.5,False,WD_ALIGN_PARAGRAPH.CENTER,8); para(d,'适用范围：Demo 0.2 至 Demo 0.4 工程开发、测试、验收与汇报',10.5,False,WD_ALIGN_PARAGRAPH.CENTER,55); para(d,'本文件定义平台的首批可交付范围、技术边界和验收标准。',10.5,False,WD_ALIGN_PARAGRAPH.CENTER,5); d.add_page_break()
head(d,'文档变更记录'); table(d,['日期','版本','变更编号','修改章节','变更说明','作者'],[['2026-09-07','V1.1','CR-001','全篇','补充需求调研、架构图、业务流程、接口边界、非功能需求与工程目录','项目组']]); head(d,'目录'); [para(d,x,10.5,False,None,2) for x in ['1 引言','2 总体描述','3 具体需求','4 非功能性需求','5 接口设计','6 设计约束','7 软件系统属性','8 需求分级','9 附录']]; d.add_page_break()
head(d,'1 引言'); head(d,'1.1 目的',2); para(d,'本文档定义智能互动教学平台的业务范围、用户角色、功能需求、非功能需求、接口边界、数据结构和版本验收标准，为产品设计、前后端开发、AI 服务开发、测试和项目验收提供共同依据。读者包括项目经理、产品经理、前端工程师、Spring Boot 后端工程师、FastAPI AI 工程师、测试人员和指导老师。'); head(d,'1.2 范围',2); para(d,'本 SRS 适用于 Demo 0.2 至 Demo 0.4。范围覆盖教师工作台、学生互动课堂、课程与课堂内容管理、AI 课堂生成、AI 教师问答、测验反馈、知识库基础检索、教师审核、结果查看和基础导出。TTS、ASR、复杂知识图谱、3D 仿真、多人实时课堂和大规模模型训练属于后续规划。'); head(d,'1.3 术语与缩略语',2); table(d,['术语','英文或缩写','说明'],[['SRS','Software Requirement Specification','软件需求规格说明书'],['AI','Artificial Intelligence','人工智能'],['LLM','Large Language Model','大语言模型'],['RAG','Retrieval Augmented Generation','检索增强生成'],['API','Application Programming Interface','应用程序编程接口'],['JWT','JSON Web Token','身份令牌']])
head(d,'2 总体描述'); head(d,'2.1 产品前景',2); para(d,'平台将备课、课堂生成、互动学习、测验反馈和学情沉淀连接成闭环。现有 Demo 0.1 已验证“主题输入 → 课堂生成 → 互动课堂 → 测验反馈”的概念链路，后续工程开发需要补齐持久化、真实模型调用、教师审核和知识库检索。'); head(d,'2.2 用户特征',2); table(d,['角色','主要场景','核心诉求'],[['教师','上传资料、生成课堂、编辑审核、发布课堂','降低备课成本，生成内容可控可改'],['学生','进入课堂、提问、答题、查看反馈','互动学习，反馈及时易懂'],['管理员','用户、配置、日志和模型状态管理','权限清晰，关键操作留痕'],['项目组','分工开发、联调、测试和汇报','接口清晰，版本可回退']]); head(d,'2.3 约束条件和依赖',2); para(d,'支持主流现代浏览器和响应式页面；业务数据需要持久化并可备份；AI 外部服务故障不能阻断演示；上传资料限制类型和大小；敏感配置通过环境变量注入；课程、学生和答题数据按角色授权访问。'); head(d,'2.4 系统架构',2); para(d,'前端采用 Vue 3；Spring Boot 业务后端负责认证、权限、课程、场景、审核、测验和事务；FastAPI AI 服务负责 LangChain 编排、RAG 检索、结构化生成、问答和 YOLOv8 推理适配；MySQL 保存业务数据，向量库保存检索索引。'); image(d,DOCS/'architecture'/'系统架构图.svg',6.65); para(d,'图 2.4.1 系统总体架构图',9,False,WD_ALIGN_PARAGRAPH.CENTER,8)
head(d,'3 具体需求'); head(d,'3.1 系统功能总览',2); para(d,'平台按教师、学生和管理员三类角色组织功能。教师侧重点是资料、课程、课堂生成与审核；学生侧重点是互动课堂、问答、测验与结果；管理员侧重点是用户、配置和审计。'); table(d,['编号','功能','优先级','首版验收标准'],[['FR-001','身份认证','P0','登录返回令牌和角色，连续错误触发限制'],['FR-002','课程 CRUD','P0','课程增删改查，刷新后数据不丢失'],['FR-003','课堂结构生成','P0','不同主题生成大纲和不少于 3 个场景'],['FR-004','教师审核','P0','内容可编辑、通过或退回，状态可追踪'],['FR-005','互动课堂','P0','按场景学习并保存进度'],['FR-006','AI 教师问答','P0','回答与课程相关，异常有明确提示'],['FR-007','测验与反馈','P0','提交后返回得分、解析和建议'],['FR-008','知识库','P0','TXT、Markdown、PDF 可上传、解析、检索'],['FR-009','知识关系','P1','查看知识点及包含、前置、相关关系'],['FR-010','学习结果','P1','展示得分、错题和知识点掌握情况'],['FR-011','导出分享','P1','至少支持 HTML 或 JSON 一种导出'],['FR-012','图像识别','P2','提供 YOLOv8 推理接口占位']]); head(d,'3.2 核心业务流程',2); image(d,DOCS/'architecture'/'课堂生成活动图.svg',5.5); para(d,'图 3.2.1 教师生成并发布课堂活动图',9,False,WD_ALIGN_PARAGRAPH.CENTER,8); image(d,DOCS/'architecture'/'平台用例图.svg',6.65); para(d,'图 3.2.2 平台核心用例图',9,False,WD_ALIGN_PARAGRAPH.CENTER,8)
head(d,'3.3 身份认证与账号管理',2); para(d,'支持学生、教师和管理员登录、注册、退出、令牌刷新和角色权限控制。账号长度 6 至 50 个字符，密码长度 8 至 32 个字符，连续 5 次密码错误后锁定 15 分钟。密码强哈希加盐保存，验证码默认 5 分钟失效，登录日志至少保留 180 天。'); head(d,'3.4 课程与课堂管理',2); para(d,'教师可以创建课程草稿，填写主题、学科、年级、课时、教学目标并关联资料。课程状态为 DRAFT、GENERATING、PENDING_REVIEW、PUBLISHED、COMPLETED 或 FAILED。教师可以调整标题、目标、场景顺序和题目内容，审核通过后方可发布。'); head(d,'3.5 AI 课堂生成与问答',2); para(d,'业务后端将主题、教学目标、当前场景和资料 ID 发送给 AI 服务。AI 服务先检索资料片段，再通过 LangChain 角色提示词生成结构化 JSON。返回结果必须通过 Pydantic 或 JSON Schema 校验。AI 教师问答保留最近若干轮上下文，并采用启发式回答。'); head(d,'3.6 测验、评分与反馈',2); para(d,'题目支持单选、多选和简答。客观题由规则评分；简答题预留 AI 辅助评价接口。提交后返回得分、正确答案、解析、关联知识点和复习建议，并保存用户、课程、题目、答案、得分和时间。'); head(d,'3.7 知识库与基础 RAG',2); para(d,'首版支持 TXT、Markdown 和 PDF 文本资料。文件上传后保存文件信息和课程关联，AI 服务完成文本清洗、分段、向量化和召回。回答返回来源资料 ID、片段序号和摘要。'); head(d,'3.8 数据存储与关系',2); image(d,DOCS/'architecture'/'数据ER图.svg',6.65); para(d,'图 3.8.1 核心数据 E R 关系图',9,False,WD_ALIGN_PARAGRAPH.CENTER,8); table(d,['实体','关键字段','用途'],[['user','id、username、password_hash、role、status','身份与权限'],['course','id、owner_id、title、subject、grade、status','课程与版本状态'],['lesson_scene','id、course_id、scene_type、content_json、sort_order','课堂场景内容'],['material','id、course_id、file_name、file_type、parse_status','课程资料与解析状态'],['question','id、course_id、scene_id、stem、answer、analysis','测验题目'],['answer_record','id、question_id、user_id、score、feedback','答题与评分结果'],['knowledge_point','id、course_id、name、parent_id、relation','知识点及关系'],['chat_message','id、course_id、user_id、role、content','问答历史']])
head(d,'4 非功能性需求'); head(d,'4.1 时间与性能',2); para(d,'普通业务接口 P95 响应时间不高于 2 秒；AI 生成和问答目标响应时间不高于 8 秒，超过阈值显示进度、超时和重试入口；课堂互动至少支持 500 人同时参与的设计目标；核心链路连续执行 3 次不得出现阻断性错误。'); head(d,'4.2 可靠性与可恢复性',2); para(d,'关键业务数据每日备份并至少保留 30 天。模型不可用时切换 mock 模式；网络失败、格式错误和资料解析失败返回统一错误码，不覆盖已有有效课程版本。'); head(d,'4.3 安全与隐私',2); para(d,'密码、令牌和 API Key 不得明文写入代码或日志；生产通信使用 HTTPS；业务接口按角色授权；上传文件限制大小和类型；删除、发布和审核记录审计日志。'); head(d,'4.4 易用性与兼容性',2); para(d,'页面采用响应式布局，关键按钮提供加载、成功和失败反馈；支持 Chrome、Edge、Firefox、Safari 等主流现代浏览器；桌面和移动宽度下不得出现主要内容横向溢出。')
head(d,'5 接口设计'); head(d,'5.1 用户接口',2); para(d,'教师端包含工作台、课程生成、课堂审核、知识库和学情结果页；学生端包含课程列表、互动课堂、问答、测验和学习结果页；管理员端包含用户管理、系统配置、模型状态和操作审计页。'); head(d,'5.2 业务接口',2); table(d,['方法','路径','用途','角色'],[['POST','/api/v1/auth/login','登录','全部'],['GET','/api/v1/courses','课程列表','教师/学生'],['POST','/api/v1/courses','创建课程草稿','教师'],['GET','/api/v1/courses/{id}','课程详情','授权用户'],['PUT','/api/v1/courses/{id}','修改课程','教师'],['DELETE','/api/v1/courses/{id}','删除课程','教师/管理员'],['POST','/api/v1/courses/{id}/generate','生成课堂','教师'],['POST','/api/v1/courses/{id}/review','提交或处理审核','教师'],['POST','/api/v1/questions/{id}/answers','提交答案','学生'],['POST','/api/v1/courses/{id}/chat','AI 教师问答','学生/教师'],['POST','/api/v1/courses/{id}/materials','上传资料','教师']]); head(d,'5.3 AI 服务接口',2); table(d,['方法','路径','输入','输出'],[['GET','/api/v1/health','无','服务状态'],['POST','/api/v1/lessons/generate','topic、subject、grade、material_ids','title、objectives、scenes、provider'],['POST','/api/v1/chat','course_context、scene、question','answer、sources、provider'],['POST','/api/v1/quiz/generate','course_id、knowledge_points、count','questions'],['POST','/api/v1/vision/detect','image_url 或文件','objects、confidence、model']]); head(d,'5.4 通信约定',2); para(d,'接口使用 REST JSON，必要时使用 WebSocket 支持流式问答或课堂状态。统一响应包含 code、message 和 data；业务后端与 AI 服务之间使用服务级鉴权和超时控制；模型结果先结构校验，再写入业务数据库。')
head(d,'6 设计约束'); head(d,'6.1 标准约束',2); para(d,'开发遵循 GB/T 8567-2006、IEEE 830-1998、RESTful API、组件化前端和分层后端规范。接口、字段和状态命名在联调前冻结，需求变更通过变更记录管理。'); head(d,'6.2 技术约束',2); para(d,'前端采用 Vue 3；业务后端采用 Java 17 与 Spring Boot 3；AI 服务采用 Python 3.11 与 FastAPI；关系数据采用 MySQL 8；向量检索采用 pgvector 或 Chroma；YOLOv8、PyTorch 和 TensorFlow 以适配器或独立训练目录接入。'); head(d,'6.3 硬件与部署约束',2); para(d,'开发演示环境建议至少 4 核 CPU、8 GB 内存、100 GB 磁盘和稳定网络。模型服务可使用外部 API 或本地推理节点；GPU 不可用时，图像识别和训练能力降级为接口占位。')
head(d,'7 软件系统属性'); head(d,'7.1 可维护性',2); para(d,'前端、业务后端和 AI 服务分目录管理；配置通过环境变量注入；统一日志、错误码和 DTO；大模型提供商、向量库和视觉模型通过适配器替换。'); head(d,'7.2 可测试性',2); para(d,'核心验收路径为：登录 → 创建主题 → 生成课堂 → 教师审核 → 进入互动课堂 → 提交测验 → 查看反馈 → 查看结果。测试覆盖空主题、超长主题、重复点击、模型超时、格式错误、网络失败、资料解析失败和越权访问。'); head(d,'7.3 可扩展性',2); para(d,'课堂场景以 scene_type 和 content_json 表达，允许后续加入白板、实验、PPTX 和项目式学习。AI 服务保留文本、语音、图像和训练模型扩展点。')
head(d,'8 需求分级'); table(d,['等级','定义','本项目示例'],[['P0','无法完成核心演示则必须实现','登录、课程、生成、审核、课堂、测验、知识库'],['P1','影响完整度但可在核心闭环后补齐','知识关系、结果页、导出、日志监控'],['P2','规划性或高风险能力，不阻断 Demo 0.4','TTS、ASR、YOLO 实际模型、3D、多用户实时课堂']]); para(d,'版本门槛：第 5 天完成 Demo 0.2；第 10 天完成 Demo 0.3；第 15 天冻结 Demo 0.4 核心功能；最后五天只处理测试、稳定性、部署和汇报材料。')
head(d,'9 附录'); head(d,'附录 A 工程目录',2); para(d,'工程骨架位于 platform/：frontend 为 Vue 3 前端，backend 为 Spring Boot 业务后端，ai-service 为 FastAPI AI 服务，docs 保存接口和开发约定。现有 Demo 0.1 页面保留在 outputs/，作为交互验收参考。'); head(d,'附录 B 需求追踪矩阵',2); table(d,['需求编号','需求名称','对应模块','验收材料'],[['FR-001','身份认证','backend / auth','登录接口测试记录'],['FR-002','课程 CRUD','backend / course','课程增删改查记录'],['FR-003','课堂结构生成','ai-service / lessons','生成 JSON 与页面截图'],['FR-004','教师审核','review module','状态流转记录'],['FR-005','互动课堂','frontend lesson','演示脚本'],['FR-006','AI 教师问答','ai-service / chat','问答测试样例'],['FR-007','测验与反馈','quiz module','答题与评分记录'],['FR-008','知识库','ai-service / rag','资料解析与检索记录'],['FR-009','知识关系','knowledge module','知识关系页面'],['FR-010','学习结果','result module','结果页截图'],['FR-011','导出分享','export module','HTML 或 JSON 导出包'],['FR-012','图像识别','ai-service / vision','YOLOv8 接口占位说明']]); head(d,'附录 C 已知边界',2); para(d,'Demo 0.1 的语音、白板、图像和部分智能体状态为交互演示，不应表述为已完成真实模型能力。真实模型接入、资料检索和结构化输出必须以接口日志或测试记录为依据。')
d.core_properties.title='智能互动教学平台软件需求规格说明书'; d.core_properties.subject='需求评审版'; d.core_properties.author='项目组'; d.save(OUT); print(OUT)
