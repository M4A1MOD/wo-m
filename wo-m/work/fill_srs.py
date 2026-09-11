from docx import Document
from pathlib import Path

src = Path(r"C:\Users\d1828\Documents\xwechat_files\wxid_i0r32xl0pplw22_3589\msg\file\2026-09\01_智能互动教学平台_Software Requirement Specification_V1.0.docx")
out = Path(r"C:\Users\d1828\Documents\Codex\2026-09-04\wo-m\outputs\01_智能互动教学平台_Software Requirement Specification_V1.0_已填写子功能.docx")
doc = Document(src)

repls = {
    '（此处填写）': '子功能键：用户名/手机号注册、密码登录、验证码校验、忘记密码、退出登录、登录状态保持。',
    '此处描述子功能中包含的功能。用例图和图。': '系统流程：用户输入账号和密码后提交；系统校验账号格式、密码完整性和验证码；校验通过后查询用户信息并生成登录会话；校验失败时返回明确错误原因。注册流程需校验账号唯一性和密码复杂度；忘记密码流程通过短信或邮箱验证码完成身份验证后重置密码。',
    '详细列出与本功能相关的功能需求。包括：该功能的预期的输入/输出，是否满足其有效性输入。还应该包括实际情况详细说明并保证必要的。 包括必要的信息，确认数据时使用。': '功能：支持学生、教师和管理员登录注册。输入为用户名/手机号、密码、验证码和用户角色；输出为登录令牌、用户基本信息和角色权限。账号长度为 6-50 个字符，密码长度为 8-32 个字符，手机号须符合中国大陆手机号格式。',
    '子功能应包含以下内容：': '流程：账号注册时，系统校验必填项、账号唯一性、密码复杂度和验证码有效期；登录时校验账号、密码、账号状态和权限；连续 5 次密码错误后锁定账号 15 分钟。',
    '子功能应包含以下内容，以及执行的操作和每个操作的过程。例如：': '输出：成功时返回统一 JSON 响应、用户 ID、角色、登录令牌和过期时间；失败时返回错误码和提示信息。网络通信失败、验证码过期、账号不存在、密码错误和账号锁定等情况均需给出可理解提示，不得返回密码等敏感信息。',
    '子功能应包含以下内容：\nA. 对该功能所涉及的数据的详细描述，包括：': '存储：用户信息保存于 user 表，登录会话保存于 session/token 存储。密码必须使用强哈希算法加盐保存，不得明文存储；验证码仅保存摘要、用途和过期时间，默认 5 分钟失效；登录日志保存登录时间、IP、设备信息和结果，至少保留 180 天。',
    '子功能应包含数据来源、数据单位、时间要求、操作场合和精度的有效输入范围等内容。': '流程：页面优化数据来源于用户操作、浏览器视口和课堂状态。系统根据设备宽度调整导航、卡片和课堂区域布局；普通页面加载目标不超过 2 秒，按钮点击反馈不超过 300 毫秒，生成课堂等长耗时操作应显示进度状态。',
    '子功能应包含输入数据的有效性检验、按照正确格式输出、对异常情况的响应等内容。': '输出：页面应统一使用响应式布局、规范字体、颜色和间距；输入框需提供必填、长度、格式和敏感内容校验；接口失败时保留用户已填写内容并显示重试入口；在小屏设备上不得出现主要内容横向溢出。',
    '子功能应包含数据保存的位置、类型、长度单位、时间、操作场合和精度的有效输入范围、对非法值的处理等内容。': '存储：页面偏好可保存在浏览器本地存储，课程和学习数据必须保存于服务端数据库。用户主题输入最长 500 字，课堂标题最长 100 字，超出限制时前端提示并阻止提交；异常数据不得覆盖已有有效课程版本。',
}

used = {}
for p in doc.paragraphs:
    txt = p.text
    if txt in repls:
        p.text = repls[txt]
        used[txt] = used.get(txt, 0) + 1
    elif txt.startswith('详细列出与本功能相关的功能需求'):
        p.text = '功能：本子功能应明确输入、处理和输出。输入数据须完成必填项、格式、长度、权限和业务规则校验；成功时返回结构化结果，失败时返回错误码、错误信息和可重试建议。'
        used['功能模板'] = used.get('功能模板', 0) + 1
    elif txt.startswith('子功能应包含以下内容，以及执行的操作'):
        p.text = repls['子功能应包含以下内容，以及执行的操作和每个操作的过程。例如：']
        used['输出模板'] = used.get('输出模板', 0) + 1

for p in doc.paragraphs:
    if p.text.strip() == '（此处填写所有可能涉及的数据存储，对表格格式体现如下）':
        p.text = ('数据存储设计：user（用户信息：id、username、password_hash、role、status、create_time）；'
                  'course（课程信息：id、title、subject、owner_id、status、created_at）；'
                  'lesson_scene（课堂场景：id、course_id、scene_type、content_json、sort_order）；'
                  'question（题目信息：id、course_id、stem、options_json、answer、analysis、knowledge_point）；'
                  'answer_record（答题记录：id、question_id、user_id、answer、score、feedback、created_at）；'
                  'knowledge_point（知识点：id、course_id、name、description、parent_id）；'
                  'learning_log（学习日志：id、user_id、course_id、action、duration、created_at）。'
                  '所有业务表使用主键 id 和创建时间字段，涉及用户的数据应遵循最小化采集、权限控制和定期备份原则。')

if len(doc.tables) > 2 and len(doc.tables[2].rows) > 1:
    row = doc.tables[2].rows[1].cells
    vals = ['2026-09-07', 'V1.1', 'CR-001', '3.2 子功能', '补充登录注册、页面优化、账号管理及数据存储子功能', '项目组']
    for c, v in zip(row, vals):
        c.text = v

doc.save(out)
print(out)
print('replaced_count', sum(used.values()))
