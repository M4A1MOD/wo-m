from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"C:\Users\d1828\Documents\Codex\2026-09-04\wo-m\outputs\智能互动教学平台项目心得与内容.docx"

def font(run, size=11, bold=None, color=None, italic=None):
    run.font.name = '等线'; run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '等线'); run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'DengXian'); run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), 'DengXian'); run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    if italic is not None: run.italic=italic
    if color is not None: run.font.color.rgb=RGBColor(*color)

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def margins(cell):
    tcPr=cell._tc.get_or_add_tcPr(); tcMar=OxmlElement('w:tcMar')
    for side,val in [('top',90),('start',120),('bottom',90),('end',120)]:
        x=OxmlElement('w:'+side); x.set(qn('w:w'),str(val)); x.set(qn('w:type'),'dxa'); tcMar.append(x)
    tcPr.append(tcMar)

doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(.85); sec.bottom_margin=Inches(.8); sec.left_margin=Inches(.95); sec.right_margin=Inches(.95); sec.header_distance=Inches(.35); sec.footer_distance=Inches(.35)
normal=doc.styles['Normal']; normal.font.name='等线'; normal._element.rPr.rFonts.set(qn('w:eastAsia'),'等线'); normal.font.size=Pt(11); normal.paragraph_format.space_after=Pt(8); normal.paragraph_format.line_spacing=1.33; normal.paragraph_format.first_line_indent=Inches(.3)
for n,s,c,b,a in [('Heading 1',16,(46,116,181),18,9),('Heading 2',13,(46,116,181),13,6),('Heading 3',12,(31,77,120),9,4)]:
    st=doc.styles[n]; st.font.name='等线'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'等线'); st.font.size=Pt(s); st.font.bold=True; st.font.color.rgb=RGBColor(*c); st.paragraph_format.space_before=Pt(b); st.paragraph_format.space_after=Pt(a); st.paragraph_format.keep_with_next=True
hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; font(hp.add_run('智能互动教学平台 · 项目心得'),9,color=(120,130,145))
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(fp.add_run('第 '),9,color=(120,130,145)); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); fp._p.append(fld); font(fp.add_run(' 页'),9,color=(120,130,145))
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(20); p.paragraph_format.space_after=Pt(8); font(p.add_run('智能互动教学平台项目心得与内容'),24,bold=True,color=(25,55,95))
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(18); font(p.add_run('——从 Demo 0.1 到可持续学习生态的设计思考'),13,color=(85,95,110),italic=True)
meta=doc.add_table(rows=2,cols=4); meta.autofit=False
vals=[['项目名称','智能互动教学平台','版本','Demo 0.1'],['文档类型','项目心得与阶段总结','适用场景','课程设计 / 项目汇报']]
for i,row in enumerate(meta.rows):
    for j,cell in enumerate(row.cells):
        cell.width=Inches([.9,2.25,.9,2.25][j]); margins(cell); shade(cell,'F4F6F9' if j%2==0 else 'FFFFFF'); cell.text=''; pp=cell.paragraphs[0]; pp.paragraph_format.space_after=Pt(0); font(pp.add_run(vals[i][j]),10,bold=(j%2==0),color=(65,78,100) if j%2==0 else (40,45,55))

def para(t):
    p=doc.add_paragraph(); p.paragraph_format.widow_control=True; font(p.add_run(t)); return p
def bullet(t):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.left_indent=Inches(.28); p.paragraph_format.first_line_indent=Inches(-.18); p.paragraph_format.space_after=Pt(4); font(p.add_run(t)); return p

sections=[
('一、项目缘起：从“会生成内容”走向“会陪伴学习”',['在人工智能快速进入教育场景的今天，很多平台已经能够完成题目生成、文本问答、课件撰写等工作，但“能生成”并不等于“能教学”。真正有价值的学习体验，应该包含目标设定、内容组织、过程互动、即时反馈和学习迁移等完整环节。我们组选择“智能互动教学平台”作为项目，正是希望把大模型的生成能力放进真实的课堂逻辑中，让 AI 不只是一个回答问题的工具，而是能够理解课程目标、识别学习状态、组织互动活动并陪伴学生成长的学习伙伴。','附件案例对本项目提出了较为完整的方向：平台需要支持课程知识库与知识图谱管理，能够根据学习主题自动生成结构化大纲和多媒体课堂；同时还要具备 AI 教师、AI 同学等多智能体协作能力，支持幻灯片、测验、交互模拟、项目式学习、白板讲解，以及 TTS/ASR 等语音能力。我们在理解这些要求时，特别注意区分“案例资料中的功能描述”和“用户真正提出的任务”。案例文档提供的是项目背景、应用设想和技术参考，并不是必须逐条照搬的指令；用户的核心请求是围绕所选项目生成一个可演示的 Demo 0.1。因此，我们把复杂愿景收敛为一个可运行、可讲解、可继续迭代的最小闭环。','这个判断对项目推进非常重要。若一开始就试图实现完整的多模型接入、真实语音链路、知识图谱可视化、PPT 导出和课堂数据分析，项目很容易陷入“功能清单很长，但没有一条流程真正跑通”的困境。Demo 0.1 的价值不在于功能数量，而在于验证核心假设：用户输入一个主题后，平台能否快速生成一节结构清晰的课；学生能否在课堂中与 AI 互动；互动结果能否转化为可理解的反馈。']),
('二、项目定位：让课堂从单向传递变成共同探索',['我们对智能互动教学平台的定位是“面向教师和学生的 AI 课堂工作台”。教师端关注备课效率和教学设计质量，学生端关注理解、练习和反馈，平台本身则负责把课程知识、智能体能力和互动场景连接起来。这个定位包含三层含义。第一，平台不是简单的聊天窗口，而是有课程结构、有场景编排、有教学节奏的产品。第二，平台不是完全替代教师，而是把教师从重复劳动中释放出来，让教师把更多时间用于启发、判断和情感支持。第三，平台不是追求让学生“看起来很热闹”，而是要让互动服务于学习目标，最终帮助学生理解知识并能够迁移运用。','在 Demo 0.1 中，我们选择“牛顿第二定律”作为示例课堂主题。这个主题既有明确的知识关系，又适合用生活情境和交互实验来展示。首页通过“最近课堂”和“平台能力”建立产品认知；生成页让用户输入学习主题，并展示 AI 教师、内容设计师、测验助教和知识校验员四类角色；课堂页则将引入、概念讲解、交互实验和即时测验组织成连续流程。这样的安排对应了真实教学中的“导入—理解—练习—检验”，也便于在演示时清楚说明系统如何工作。']),
('三、Demo 0.1 的核心内容与实现思路',['工作台是用户进入平台后的第一印象。我们没有把页面设计成复杂的数据驾驶舱，而是采用清晰的卡片布局，突出最近课堂、学习进度和平台能力。最近课堂包含进行中、已完成和草稿三种状态，分别代表学习过程、结果沉淀和内容准备。这样的状态设计虽然是模拟数据，但已经体现出平台未来的核心对象：课程不是一次性生成后就结束，而是会经历创建、学习、复习、评估和迭代。','生成课堂页是 Demo 0.1 的核心入口。用户可以输入“牛顿第二定律”等主题，也可以点击快捷主题。点击生成后，系统会展示一套四步课堂结构：生活情境引入、概念讲解、交互实验和即时测验。虽然当前版本使用的是本地模拟数据，但交互过程完整呈现了未来后端服务应承担的职责：主题解析、知识点抽取、教学目标匹配、场景编排和测验生成。','在产品设计上，结构化大纲比一段长篇生成文本更适合教学。大纲可以让教师快速浏览并修改，也可以让系统进一步把每个条目转换成幻灯片、白板或实验场景。例如“概念：力、质量与加速度”可以生成公式推导和动画演示；“实验：改变变量，观察运动”可以生成质量、受力和加速度三个滑块；“检验：3 题即时测验”则可以生成单选题、简答题以及分层评价标准。可见，大纲不是最终内容，而是不同教学场景之间的中间层。','案例文档中强调了多智能体技术。我们在 Demo 中把它转化为四个易于理解的角色。AI 教师负责讲解、追问和课堂节奏；内容设计师负责组织大纲和互动场景；测验助教负责题目生成和评价；知识校验员负责检查课堂内容是否与知识库和课程目标一致。实际系统中，这些角色可以由同一个大语言模型通过不同提示词和工具调用实现，也可以接入不同模型服务。重要的是把职责拆开，让生成过程可解释、可调试、可替换。','进入互动课堂后，学生首先看到的是一个生活化问题：“为什么推车会越推越快？”这个问题没有直接给出公式，而是先引导学生调用已有经验。随后页面展示力 F、质量 m、加速度 a 和 F = ma 等知识标签，将抽象概念逐步聚焦。AI 教师提示语采用追问方式，而不是把答案一次性说完，这体现了支架式教学的思路。课堂页还提供播放讲解、向老师提问和白板标注三个操作。即时测验采用单选题，学生选择“B. 加倍”后会看到基于公式的解释，并获得学习积分。']),
('四、技术与产品设计方面的收获',['这次实践让我更加理解 MVP（最小可行产品）的意义。MVP 不是随便做一个简陋页面，而是选择最能验证产品价值的一条路径。对本项目而言，这条路径就是“主题输入—课堂生成—互动学习—结果反馈”。只要这条链路顺畅，后续接入真实模型、数据库和多媒体能力就有清晰的落点；如果链路本身不成立，继续增加功能只会放大问题。','教育场景对准确性、稳定性和可解释性的要求高于普通内容应用。大模型可能产生事实错误、概念混淆或超出课程范围的内容，因此知识库和知识图谱不能只是宣传概念，而要真正参与生成控制。未来的系统应当为课程设定知识边界、年级难度、教学目标和参考资料，并在生成后给出来源、置信度或校验结果。','在设计课堂场景时，我曾经容易被“3D 模拟、动画、语音、小游戏”等功能吸引，但后来意识到，互动形式必须服从学习目标。比如牛顿第二定律的重点是理解力、质量和加速度的关系，那么滑块实验、变量对比和公式推导比单纯的动画更有效。未来做功能取舍时，应先问“学生要学会什么”，再问“什么互动最能帮助他学会”。']),
('五、团队协作与项目推进中的体会',['智能互动教学平台横跨产品、前端、后端、人工智能和教育设计多个领域，单靠一个人很难同时做到技术可行和教学合理。项目推进中，团队需要尽早统一三个问题：目标用户是谁；第一版要验证什么；如何定义完成。我们把 Demo 0.1 的验收标准定义为页面能够打开、主题能够生成、课堂能够进入、题目能够作答、反馈能够显示，这使得协作有了共同尺度。','我还认识到，团队沟通不能只交流“做没做完”，更要交流“为什么这样做”。例如把知识校验员作为独立角色，并不是为了让页面看起来更智能，而是为了回应教育场景对可信度的要求；把课堂分成四个场景，也不是为了增加步骤，而是为了让教学目标、内容和评价形成闭环。']),
('六、当前不足与问题反思',['当前版本主要使用本地模拟数据，尚未接入真实的大语言模型、数据库和用户账号体系，生成结果还不能真正根据资料动态变化。','课堂内容以文字和卡片为主，幻灯片、白板、交互实验、TTS/ASR 等多模态能力还停留在界面预留和交互模拟阶段。','测验类型目前以单选题为主，尚未覆盖多选、简答、编程题和项目式任务，也没有形成完整的学习画像。','知识库和知识图谱尚未实现真实的上传、解析、实体抽取、关系编辑和可视化，知识一致性校验仍需进一步工程化。','平台还缺少教师审核、课堂管理、权限控制、隐私保护和数据脱敏等机制，距离真实部署仍有较长距离。']),
('七、后续迭代规划',['第一阶段是“真实生成”。接入 DeepSeek、通义千问或其他模型服务，建立统一的模型适配层，实现主题解析、课程大纲生成、题目生成和问答服务，并保存每次生成的版本。','第二阶段是“知识增强”。建立课程知识库，支持 PDF、PPTX、Word 和网页资料解析；通过切分、向量检索和知识图谱构建，为生成过程提供可引用的课程依据。','第三阶段是“场景丰富”。把大纲条目转换为可编辑幻灯片、互动白板、模拟实验和知识小游戏；接入 TTS 和 ASR，实现 AI 语音授课、学生口述回答和课堂转写。','第四阶段是“学习评价”。记录学生的答题过程、提问内容、停留时间和知识点掌握情况，形成可解释的学习报告。同时建设隐私保护、权限管理和数据脱敏机制。']),
('八、结语：技术最终要回到人的成长',['通过这次智能互动教学平台项目，我对人工智能应用开发的理解发生了变化。以前我更容易把注意力放在模型能力、页面效果和功能数量上，认为只要生成得足够快、界面足够丰富，产品就会有吸引力。现在我更加相信，教育产品的核心不是“替学生完成学习”，而是帮助学生更好地发现问题、理解问题、表达问题并解决问题。AI 可以承担解释、示范、陪练和反馈，但学习的主动性、教师的判断力以及人与人之间的信任，仍然是课堂不可替代的部分。','Demo 0.1 只是一个起点，它还不是真正意义上的完整平台，却已经把一个抽象的选题变成了可操作、可演示、可讨论的产品雏形。更重要的是，它让我们看见了从概念到落地所需要经历的过程：理解需求、收敛范围、设计闭环、实现交互、验证路径、暴露问题，再通过迭代逐步接近真实场景。未来无论是继续完善这个项目，还是参与其他人工智能应用开发，我都会坚持以用户目标为起点，以可验证的闭环为抓手，以安全、可信和有教育价值为底线，让技术真正服务于人的成长。'])]

for title,ps in sections:
    doc.add_heading(title, level=1)
    for t in ps: para(t)
    if title.startswith('六、当前不足'):
        for t in ['当前版本主要使用本地模拟数据，尚未接入真实模型、数据库和用户体系。','多媒体课堂、白板、交互实验、TTS/ASR 仍是待实现能力。','测验类型和学习画像还不够丰富。','知识库、知识图谱和审核机制需要进一步工程化。','真实部署还需补充权限、隐私和数据安全设计。']: bullet(t)
doc.save(OUT); print(OUT)
