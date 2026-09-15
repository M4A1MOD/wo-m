from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

src = r'C:\Users\d1828\Downloads\学号_姓名_软件工程专业见习周志-第1周.docx'
out = r'C:\Users\d1828\Documents\Codex\2026-09-04\wo-m\学号_姓名_软件工程专业见习周志-第1周-Demo0.2.docx'
doc = Document(src)
table = doc.tables[0]
table.cell(1, 1).text = '第1周'
table.cell(1, 3).text = '2026年9月7日--2026年9月11日'
table.cell(2, 1).text = '董爵爷'
table.cell(2, 3).text = '202424120229'
table.cell(2, 5).text = '软工2班'
table.cell(3, 2).text = '信息楼317教室'
content = [
('周一（9月7日）', '熟悉智能互动教学平台 Demo 0.1 原型和需求文档，梳理 Demo 0.2 第一周计划。明确课程、课堂场景、题目和答题记录的数据关系，确定本周以课程 CRUD、课堂结构生成和本地持久化为主要目标。'),
('周二（9月8日）', '完成 Demo 0.2 基础工程搭建。前端采用 Vue + Vite，配置 /api 代理；后端采用 Spring Boot，建立 Course 实体、Repository、Controller 和 H2 数据库配置；AI 服务使用 FastAPI 提供结构化课堂生成接口。'),
('周三（9月9日）', '实现课程新增、查询、修改和删除接口，并在启动时写入“牛顿第二定律”示例课程。前端生成课堂后保存草稿，工作台动态展示课程列表，课程状态、场景数量和更新时间可以回显。'),
('周四（9月10日）', '完成前后端联调和异常回退。后端可用时通过 API 持久化课程；服务不可用时自动使用 localStorage，刷新页面后仍能恢复课程。互动课堂保留讲解、提问和随堂测验等 Demo 0.1 交互，并补充 Demo 0.2 的提示反馈。'),
('周五（9月11日）', '完成 Demo 0.2 验收：验证课程生成、刷新恢复、课程删除、互动课堂入口和 AI mock 结构化输出；将默认入口切换为 Demo 0.2，并删除旧 Demo 0.1 页面文件。总结本周问题，下一阶段继续完善场景、题目和答题记录模型。')]
body = table.cell(4, 1)
body.text = ''
for day, text in content:
    p = body.add_paragraph()
    p.add_run(day + '：').bold = True
    p.add_run(text)
reflection = table.cell(5, 1)
reflection.text = '本周围绕 Demo 0.2 完成了从需求梳理到可运行原型的迭代。通过课程 CRUD、后端 H2 持久化、AI mock 接口和 localStorage 降级，平台已经形成基本闭环。联调过程中认识到前后端接口字段需要尽早统一，异常回退也应在原型阶段设计。下一周将继续细化课堂场景、题目和答题记录，并逐步接入更真实的智能生成能力。'
for row in table.rows:
    for cell in row.cells:
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.name = '宋体'; run.font.size = Pt(10.5)
doc.save(out)
print(out)
