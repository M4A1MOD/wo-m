from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "学号_姓名_软件工程专业见习周志-第3周-Demo0.3.docx"

content_paragraphs = [
"周一 9月14日 版本梳理与整合方案设计。本周工作的核心是把前一阶段形成的 Demo 0.2 基础能力与新开发的 Demo 0.3 多学科课堂能力整合为同一个正式版本。首先，我重新检查了前端 Vue 工程、Spring Boot 业务后端、FastAPI 智能服务以及独立 HTML 演示文件之间的关系，重点确认课程新增、查询、修改、删除、审核状态和本地持久化没有因为版本升级而失效。随后，我对旧版本的数据结构进行了比对。Demo 0.2 主要保存课程标题、学科、年级、状态和更新时间，而 Demo 0.3 还需要教学目标、教学重点、教学难点、课堂场景、测验、作业、资源及评价方案。为了避免旧数据升级后出现空白页面，我设计了数据补全函数：读取旧课程后，按照课程主题和学科生成默认课堂结构，再用旧字段覆盖对应内容。对于重复编号、缺失数组和异常状态也增加了兼容处理。当天还统一了版本入口，决定让工程版与单文件演示版使用同一套 Vue 源代码，独立 HTML 只负责内联构建产物，不再单独维护一份简化页面。这样既保留了 Demo 0.2 已经验证过的课程管理闭环，也为 Demo 0.3 后续增加多学科生成、课堂问答和导出功能建立了稳定基础。",
"周二 9月15日 九学科差异化内容完善。此前原型对物理课堂的描述相对完整，但其他学科仍然存在内容通用、活动形式重复的问题。为此，我按照不同学科的核心思维方式重新设计语文、数学、英语、历史、地理、生物、化学、道德与法治和物理九类模板。语文课堂突出文本细读、语言品味和证据表达，通过圈点批注、关键词替换、观点证据链等活动引导学生从具体语句形成解释；数学课堂突出概念建构、数形结合和变式应用，用参数实验、正反例判断与一题多解帮助学生理解条件变化；英语课堂围绕听说输入、功能句型和真实交际，设置预测、信息差、角色扮演及口语评价。历史课堂加入时间轴、地图定位、史料来源判断、多因素因果链和历史解释，避免只呈现结论；地理课堂围绕位置、分布、图表数据、自然与人文要素联系以及区域决策组织任务。生物、化学和物理保留科学探究共同框架，但在结构功能、宏观微观符号、变量关系等方面使用不同的知识表达。道德与法治则采用生活案例、规则辨析、议题讨论和公民行动方案。每门学科都形成教学目标、重难点、五个课堂场景、师生活动、互动任务、测验解析、作业、资源与评价标准，生成结果不再只是替换课程标题。",
"周三 9月16日 AI课堂生成与上下文问答联调。在内容模板稳定后，我继续完善前端、业务后端和智能服务之间的调用链。前端生成课堂时提交主题、学科、年级和课时，Spring Boot 通过统一接口转发给 FastAPI，智能服务根据学科规则返回结构化课程数据。为了保证演示稳定，我没有把页面完全依赖于外部模型，而是保留了按学科生成的本地降级方案：智能服务在线时优先合并接口结果，服务不可用时仍可从前端模板生成完整课堂，并明确显示当前运行状态。课堂问答部分也采用相同思路。学生在某个课堂场景中提问时，系统会携带课程主题、学科、年级、当前场景、知识点和最近对话记录，使回答能够围绕正在学习的内容，而不是给出脱离课程的一般性答案。默认回答采用启发式方式，提示学生先定位关键信息、建立知识关系、寻找证据，再回答一个追问，从而体现教师引导而非直接公布结论。联调过程中，我重点处理了前后端字段命名差异，例如 keyPoints 与 key_points、teacherActivity 与 teacher_activity，并通过合并函数确保服务返回不完整时仍能使用本地内容补齐。最终，生成、保存、打开、问答和重新加载形成连续流程，全栈服务和离线模式都能完成课堂演示。",
"周四 9月17日 互动课堂、测验与审核发布完善。我将生成后的课程组织为五个连续课堂场景，左侧显示场景名称、时长和互动方式，中间区域展示场景说明以及教师活动、学生活动和知识点，右侧保留教学目标、重点、难点和评价摘要。场景切换时同步更新进度，并清理上一场景的临时答题状态，避免答案显示错位。测验模块根据当前课程内容生成题目、选项、正确答案、知识点和解析。学生选择答案后，页面会立即区分正确项和错误项，显示原因说明，并通过提示语引导继续复习。教师管理方面，我保留了 Demo 0.2 的课程列表和删除功能，同时加入草稿与已发布状态切换。课程生成后默认进入待审核状态，教师检查课堂流程、内容准确性和测验答案后，可以一键发布；发现问题时也可以退回草稿继续调整。为解决网络不稳定造成的演示中断，我检查了 H2 数据库与 localStorage 的协同策略：后端在线时同步课程基础信息，离线时把完整课堂对象写入浏览器存储，刷新页面仍可恢复场景、测验和审核状态。当天还对语文、数学、英语、历史和地理五个典型主题逐一试用，确认不同课程的标题、活动语言、互动方法和评价方式具有明显区别，课堂导航、即时反馈及审核操作均能正常衔接。",
"周五 9月18日 课堂导出、独立版本构建与综合验收。本周最后一天重点完善课堂导出的完整性。原有导出内容偏向页面展示，缺少教师真正需要的教学设计信息，因此我重新组织 HTML 课堂文件，将课程基本信息、教学目标、教学重点和难点、逐场景的教师活动与学生活动、互动方式、知识点、测验题目、参考答案、解析、课后作业、资源清单和评价方案全部写入。导出的 HTML 文件可以脱离平台直接打开，也可以打印或另存为 PDF。与此同时，增加带 schemaVersion 和导出时间的 JSON 数据文件，便于后续系统继续读取、交换或扩展课程内容。为了让演示入口与正式工程保持一致，我使用 Vite 对 Vue 项目进行生产构建，再把生成的 JavaScript 和 CSS 内联到单个 HTML 文件中，并特别处理脚本结束标签和字符串替换可能造成的内容截断问题。综合验收按照完整用户路径进行：打开工作台，按学科筛选课程，生成新的历史或地理课堂，切换五个场景，向 AI 教师提问，完成测验，审核发布，导出 HTML 与 JSON，关闭服务后再次刷新检查离线恢复。最后运行 TypeScript 类型检查、前端生产构建和 Python 语法检查，并核对独立 HTML 不再引用外部静态资源。本周由此完成了 Demo 0.3 的统一正式版，保留旧版数据和功能，同时形成可展示、可保存、可交互、可导出的多学科课堂闭环。",
]

reflection_paragraphs = [
"一 本周最明显的收获，是理解了版本升级不等于简单增加页面。Demo 0.3 如果只把多学科模板放到 Demo 0.2 旁边，表面上功能更多，实际会出现两套入口、两种数据结构和两种交互逻辑。用户不知道应该打开哪个文件，开发者修复问题时也要修改两处。真正的整合必须先确定稳定基线，再把新能力接入已有闭环。因此，我把 Demo 0.2 的课程管理、持久化和离线恢复看作基础设施，把 Demo 0.3 的学科内容、AI 问答、测验和导出看作上层能力。统一数据模型、统一入口和统一构建流程后，版本名称才与实际交付物一致。这让我认识到软件工程中的“完成”不仅是代码可以运行，还包括版本边界清楚、旧数据能够迁移、用户路径连续、维护成本可控。以后开发新版本时，我会更早考虑迁移策略和兼容层，而不是等功能全部写完后再处理历史问题。",
"二 在完善其他学科的过程中，我进一步体会到业务建模的重要性。最初容易把“支持多学科”理解为更换学科名称和几个知识点，但真正试用后会发现，语文阅读、数学探究、英语交际、历史解释和地理决策的学习过程差异很大。如果所有课程都使用“教师讲解、学生讨论、完成练习”这样的通用句式，系统虽然能生成内容，却没有教学价值。本周我从每门学科的证据来源、思维路径、典型活动和评价方式出发设计模板。例如历史结论需要史料来源与相互印证，地理解释需要地图位置和数据特征，语文观点需要具体文本，数学结论需要条件、表示与推理。将这些规则落实到字段和场景后，生成内容才具有可辨识性。这个过程也说明，AI 教育产品不能只追求语言流畅，还要把学科知识、教学方法和课程结构转化为明确约束。大模型负责扩写和表达，系统设计负责规定边界，两者结合才能获得稳定结果。",
"三 我对前后端接口设计也有了更具体的认识。Demo 0.3 同时使用 TypeScript、Java 和 Python，同一个概念会以不同命名习惯出现。如果不提前约定字段，就会产生前端显示为空、后端验证失败或部分数据被覆盖的问题。本周处理 keyPoints 与 key_points 等差异时，我认识到接口不仅是传输数据的通道，更是各模块共同遵守的合同。一个可靠接口需要明确请求字段、可选项、默认值、响应结构、异常状态和版本号。对于 AI 返回内容，还必须考虑数组为空、字段缺失、格式不符合预期等情况，不能假设模型每次都完全正确。合并函数和类型检查虽然增加了一些代码，却把不确定性限制在边界处，使页面内部继续使用稳定的 LessonPlan 结构。以后遇到跨语言服务，我会优先设计 DTO、Schema 和示例响应，并为兼容和错误处理留下位置，这比调试时逐项猜测高效得多。",
"四 离线降级方案让我重新理解了系统可靠性。教学演示往往发生在网络环境不稳定、模型密钥未配置或某个服务临时无法启动的情况下。如果平台只能在全部组件在线时工作，功能再丰富也可能在演示现场失去价值。因此，本周继续保留 H2 与 localStorage 两层存储，并让 AI 服务不可用时回退到学科模板和启发式回答。这里的关键不是隐藏故障，而是让系统明确告诉用户当前使用哪种模式，同时保持核心操作可继续进行。在线模式提供服务端持久化与智能生成，离线模式保证创建、打开、切换、测验、审核和导出仍然有效。这种设计也帮助开发和测试，因为前端可以在没有完整环境时独立验证。我的体会是，降级不是临时补丁，而应当作为产品能力设计；需要提前划分哪些功能必须可用、哪些功能可以替代、哪些状态必须告知用户。",
"五 课堂导出的完善让我意识到，导出文件不是页面截图，而是一种新的使用场景。教师把课堂导出后，可能需要离线查看、打印备课、保存归档，也可能把结构化数据交给其他系统。如果文件只包含当前可见画面，许多隐藏在界面中的目标、答案、作业和评价就会丢失。为此，我从接收者需求重新设计导出结构：先给出课程基本信息和教学设计，再按课堂顺序列出每个场景，随后集中呈现测验答案、课后任务、资源和评价。HTML 强调阅读、打印和离线打开，JSON 强调字段完整、版本可识别和后续处理。两种格式承担不同职责，但来源必须是同一个课程对象。本周还遇到内联脚本替换时特殊字符影响构建结果的问题，这提醒我导出功能不仅要看按钮是否触发下载，还要真正打开文件，检查内容、样式、编码和资源依赖。",
"六 测试过程使我更加重视以用户路径组织验收。单独检查函数或接口能够发现局部错误，却不能保证学生或教师从开始到结束顺利完成任务。本周我按照“生成课程、保存课程、打开课堂、切换场景、提问、答题、审核、导出、刷新恢复”的顺序进行验证，并用多个学科主题检查内容差异。这样做发现了单项测试不容易暴露的问题，例如切换课程后仍保留上一题答案、后端同步基础字段时覆盖本地完整场景、独立 HTML 与工程版功能不一致等。解决这些问题后，我把验收结果整理为功能清单，明确通过项、降级行为和当前边界。这个过程让我认识到测试并非开发结束后的简单确认，而是帮助重新审视数据流和状态流的重要方法。下一阶段应继续增加自动化测试，用固定输入验证九学科结构、旧数据迁移和导出字段，降低人工回归成本。",
"七 从个人能力上看，本周的工作把前端界面、后端业务、AI 服务和构建交付连接起来。以前我更关注某一段代码是否实现，现在开始从系统角度考虑数据从哪里产生、经过哪些服务、失败后怎样处理、最终如何交付给用户。在 Vue 部分，我练习了响应式状态、计算属性、课程筛选、条件渲染和复杂页面交互；在 Spring Boot 部分，我理解了控制器代理、DTO 校验和持久化边界；在 FastAPI 部分，我使用 Pydantic 约束输入输出，并把学科生成逻辑封装成可替换的服务。构建独立 HTML 时，还需要理解 Vite 产物、资源路径和脚本内联。这些任务看似分散，实际共同服务于一条完整业务流程。它们让我更加明确，全栈开发的价值不是会使用更多框架，而是能够在模块之间建立稳定联系，并在出现问题时快速定位属于界面、数据、接口还是构建环节。",
"八 本周也暴露了一些仍需改进的地方。首先，Demo 0.3 的智能生成主要依靠学科规则和稳定模拟结果，真实模型接入后还需要处理提示词版本、响应校验、调用成本和超时重试。其次，当前后端只持久化课程基础信息，完整课堂结构主要保存在浏览器中，若要支持多设备和多人协作，应把场景、测验、资源、答题记录拆分为正式实体并设计关联关系。再次，测验数量和题型还较少，学情分析目前只体现即时反馈，尚未形成长期掌握度模型。语音、白板和真正的多媒体素材生成也不属于本周完成范围，不能因为界面中出现 AI 字样就夸大实现程度。下一步我计划先稳定课程结构的服务端存储，再扩展题型、学习记录和知识点关联；同时建立真实模型与本地降级模型的统一适配接口。通过这周的实践，我更加重视如实描述版本能力，也明白清晰列出边界能够帮助团队制定下一阶段任务。",
"九 总结本周，我认为 Demo 0.3 最重要的成果不是增加了多少按钮，而是形成了一个较完整的多学科课堂闭环。教师可以选择学科和主题生成课程，检查符合学科特点的目标与活动，在互动课堂中完成问答和测验，审核后发布，并将完整教学设计导出；旧课程能够继续使用，服务异常时也有可理解的替代路径。这一结果让我感受到，软件项目每一次迭代都应围绕可验证的用户价值展开。需求分析决定做什么，领域建模决定内容是否专业，接口和数据设计决定模块能否协作，测试与导出决定成果能否真正被使用。第三周的开发让我从“把功能写出来”逐渐转向“把版本做完整”。今后的工作中，我会继续保持版本记录、功能矩阵和验收清单，先保证主流程稳定，再逐步增加复杂能力，使每个阶段都有可以独立展示和持续演进的成果。",
]

def char_count(items):
    return sum(len(x.replace(" ", "")) for x in items)

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

def apply_font(run, name="Microsoft YaHei", size=10.5, bold=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)

def style_paragraph(p, first_line=True, after=5, line=1.45):
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.74) if first_line else None
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.widow_control = True
    for run in p.runs:
        apply_font(run)

def add_section_content(cell, paragraphs):
    cell.text = ""
    for index, text in enumerate(paragraphs):
        p = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        first, rest = text.split(" ", 1)
        label_parts = first
        if index < 5 and "日" not in first:
            # Daily labels contain spaces; take the first three tokens instead.
            tokens = text.split(" ", 3)
            label_parts = " ".join(tokens[:3])
            rest = tokens[3]
        r = p.add_run(label_parts + " ")
        apply_font(r, size=10.5, bold=True)
        r = p.add_run(rest)
        apply_font(r, size=10.5)
        style_paragraph(p, first_line=False, after=8, line=1.5)

doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.left_margin = Cm(2.4)
section.right_margin = Cm(2.4)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)
section.header_distance = Cm(1.0)
section.footer_distance = Cm(1.0)

normal = doc.styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)

title = doc.add_paragraph()
title.style = doc.styles["Title"]
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(16)
run = title.add_run("软件工程专业见习周志")
apply_font(run, size=20, bold=True)

meta = doc.add_table(rows=4, cols=6)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta.autofit = False
widths = [Cm(1.8), Cm(3.3), Cm(1.8), Cm(3.3), Cm(1.8), Cm(3.3)]
for row in meta.rows:
    for i, cell in enumerate(row.cells):
        cell.width = widths[i]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell, 90, 90, 90, 90)

values = [
    ["专 业", "软件工程", "年 级", "2024级", "周 数", "第三周"],
    ["姓 名", "________", "学 号", "202424120229", "班 级", "软件2班"],
    ["时 间", "2026年9月14日—2026年9月18日", "", "", "", ""],
    ["见习地点", "信息楼317机房", "", "", "", ""],
]
for r, row_vals in enumerate(values):
    if r in (2, 3):
        merged = meta.cell(r, 1).merge(meta.cell(r, 5))
        row_vals[1] = row_vals[1]
    for c, value in enumerate(row_vals):
        if r in (2, 3) and c > 1:
            continue
        p = meta.cell(r, c).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = p.add_run(value)
        apply_font(rr, size=10, bold=(c % 2 == 0))
        if c % 2 == 0:
            set_cell_shading(meta.cell(r, c), "EAF0F8")

doc.add_paragraph().paragraph_format.space_after = Pt(2)

def add_body_section(title_text, items):
    heading = doc.add_paragraph()
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.space_after = Pt(8)
    rr = heading.add_run(title_text)
    apply_font(rr, size=14, bold=True)
    for item in items:
        p = doc.add_paragraph()
        tokens = item.split(" ", 3)
        if title_text == "实习内容":
            label = " ".join(tokens[:3])
            rest = tokens[3]
        else:
            label = " ".join(tokens[:2])
            rest = " ".join(tokens[2:])
        r1 = p.add_run(label + " ")
        apply_font(r1, bold=True)
        r2 = p.add_run(rest)
        apply_font(r2)
        style_paragraph(p, first_line=True, after=7, line=1.5)

add_body_section("实习内容", content_paragraphs)
add_body_section("实习心得", reflection_paragraphs)

summary = doc.add_paragraph()
summary.alignment = WD_ALIGN_PARAGRAPH.RIGHT
summary.paragraph_format.space_before = Pt(8)
r = summary.add_run(f"实习内容约 {char_count(content_paragraphs)} 字    实习心得约 {char_count(reflection_paragraphs)} 字")
apply_font(r, size=9)
r.font.color.rgb = RGBColor(100, 100, 100)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("智能互动教学平台 Demo 0.3    第三周")
apply_font(run, size=8.5)
run.font.color.rgb = RGBColor(110, 110, 110)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
print("CONTENT_CHARS", char_count(content_paragraphs))
print("REFLECTION_CHARS", char_count(reflection_paragraphs))
