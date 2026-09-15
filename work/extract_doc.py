from docx import Document
p=r"C:\Users\d1828\Downloads\人工智能开发选题案例介绍.docx"
d=Document(p)
with open(r"C:\Users\d1828\Documents\Codex\2026-09-04\wo-m\work\doc_extract.txt","w",encoding="utf-8") as f:
  for i,x in enumerate(d.paragraphs):
    if x.text.strip(): f.write(f"P{i}: {x.text}\n")
  for ti,t in enumerate(d.tables):
    f.write(f"TABLE {ti}\n")
    for row in t.rows: f.write(" | ".join(c.text for c in row.cells)+"\n")
